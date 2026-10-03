"""文档解析：PDF / DOCX / Markdown -> ParsedDocument。

明确不做（边界声明，不是能力不足）
    - 不做 OCR：扫描件（无文本层）直接抛错，不静默返回空
    - 不做版面分析：双栏 PDF 的文字顺序可能错乱
    - 不做表格结构化：表格被拍平成文本行

设计要点：
    - 保留 page，引用回溯需要
    - 页眉页脚按【跨页重复】启发式剔除
    - doc_id 由文件内容哈希派生，保证同一文件重复导入得到相同 ID
"""

import hashlib
import re
from pathlib import Path
from typing import Callable

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.logger import logger
from app.core.schemas import Block, ParsedDocument


class DocumentParseError(Exception):
    """解析失败：格式不支持、文件损坏、扫描件无文本层。"""


# 扫描件判定：页数 >= N 且平均每页字符数 < M
_SCAN_MIN_PAGES = 2
_SCAN_CHARS_PER_PAGE = 50

# 页眉页脚判定：同一行出现在超过该比例页面的首行/末行
_HEADER_FOOTER_RATIO = 0.6

# ---------- 页码类页眉页脚 ----------
#
# ← Day 3 实测教训：「跨页重复」启发式只对【固定文本】有效。
# 页码每页都不同（"第 3 页" vs "第 4 页"），字符串计数恒为 1，
# 永远达不到阈值 —— 真实合同里 "第 N 页 共 49 页" 一个都没被剔掉。
# 带变量的页眉页脚（页码、日期、生成时间）必须靠【模式匹配】，不能靠字符串相等。
_RE_PAGE_FOOTER = re.compile(r"^第\s*(\d+)\s*页\s*共\s*(\d+)\s*页$")

_PAGE_NUMBER_PATTERNS = (
    _RE_PAGE_FOOTER,
    re.compile(r"^第\s*\d+\s*页$"),
    re.compile(r"^\d+\s*/\s*\d+$"),
    re.compile(r"^Page\s+\d+\s+of\s+\d+$", re.IGNORECASE),
)


def _is_page_number(line: str) -> bool:
    """整行只有页码 —— 这种行不可能是正文，可全位置删除。"""
    return any(pattern.match(line) for pattern in _PAGE_NUMBER_PATTERNS)


# Markdown 装饰符。嵌入用的文本应该是纯文本——"##"、"**" 对语义模型是纯噪声
_MD_HEADING = re.compile(r"^\s{0,3}#{1,6}\s*")
_MD_LIST = re.compile(r"^\s{0,3}[-*+]\s+")
_MD_EMPHASIS = re.compile(r"\*\*|__|`")


def _normalize(text: str) -> str:
    """空白归一化：全角半格 -> 半角，压缩连续空白，去首尾。"""
    text = text.replace("\u3000", " ").replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def _strip_markdown(text: str) -> str:
    """去掉 Markdown 装饰符，保留文字"""
    text = _MD_HEADING.sub("", text)
    text = _MD_LIST.sub("", text)
    text = _MD_EMPHASIS.sub("", text)
    return text.strip()


def _strip_headers_footers(pages: list[list[str]]) -> list[list[str]]:
    """剔除页眉页脚。两类规则并用，只做一条都会漏：

    **规则 1（模式）**：整行只是页码 —— 全位置删除。
        内容每页都不同，靠字符串计数永远识别不出来，必须用正则。

    **规则 2（重复）**：在超过阈值的页面**首行/末行**重复出现的固定文本 —— 只在首末位置删除
        （同样的字符串出现在正文中间时保留，避免误删）。
    """
    # --- 规则 1：页码行（单行即可判定，无需统计）---
    cleaned: list[list[str]] = [[line for line in lines if not _is_page_number(line)] for lines in pages]

    if len(cleaned) < 3:  # 页数太少，重复统计没有意义
        return cleaned

    # --- 规则 2：跨页重复的固定文本 ---
    threshold = max(2, round(len(cleaned) * _HEADER_FOOTER_RATIO))

    counter: dict[str, int] = {}
    for lines in cleaned:
        if not lines:
            continue

        for line in {lines[0], lines[-1]}:  # set 去重，单行页只算一次
            counter[line] = counter.get(line, 0) + 1

    repeated = {line for line, n in counter.items() if n >= threshold}
    if repeated:
        logger.debug(f"剔除疑似页眉页脚 {len(repeated)} 条，例如: {sorted(repeated)[:3]}")

    return [
        [line for i, line in enumerate(lines) if not ((i == 0 or i == len(lines) - 1) and line in repeated)]
        for lines in cleaned
    ]


def _audit_pages(pages: list[list[str]]) -> dict:
    """PDF 页级自检：页码单调性 + 空页。

    ← 调用方必须在 _strip_headers_footers **之前**调用本函数 —— 页码正是要被剔掉的对象。

    这两个检查成本极低，但能发现人工翻块列表绝对看不出来的问题。
    Day 3 实测 A 文档（51 物理页 / 文档自身 49 页）：

        extract# 1  -> (无文本)
        extract# 2  -> 49/49     ← 文档末页出现在物理第 2 页（错位）
        extract# 3  ->  1/49
        ...
        序列 [49, 1, 2, ..., 48, 49]，单调步数 48/49

    这个错位直接导致了 chunk_id 碰撞 —— 同一段末页文本在开头又出现了一次。
    """
    numbers: list[int] = []
    for lines in pages:
        for line in lines:
            match = _RE_PAGE_FOOTER.match(line)
            if match:
                numbers.append(int(match.group(1)))
                break

    empty_pages = [i + 1 for i, lines in enumerate(pages) if not lines]
    regressions = [(a, b) for a, b in zip(numbers, numbers[1:]) if b < a]

    if regressions:
        logger.warning(
            f"页脚页码非单调：{len(regressions)} 处回退 / 共 {len(numbers)} 页，"
            f"例如 {regressions[:3]} —— 物理页序与文档页码不符，文本顺序可能错位，"
            f"引用回溯的位置信息需谨慎使用"
        )
    if empty_pages:
        logger.warning(
            f"{len(empty_pages)} 页未提取到文本（第 {empty_pages[:5]} 页）—— "
            f"可能是整页图片（扫描的签字页 / 封面），这些页的内容会缺失"
        )

    return {
        "page_numbers": numbers,
        "page_order_regressions": len(regressions),
        "empty_pages": empty_pages,
    }


# ---------- 各格式解析器：统一返回 (blocks, meta) ----------


def parse_pdf(path: Path) -> tuple[list[Block], dict]:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise DocumentParseError(f"无法打开 PDF: {path.name} ({exc})") from exc

    total_pages = len(reader.pages)
    pages: list[list[str]] = []
    for page in reader.pages:
        raw = page.extract_text() or ""
        lines = [_normalize(line) for line in raw.splitlines()]
        pages.append([line for line in lines if line])

    text_chars = sum(len(line) for lines in pages for line in lines)

    # 扫描件必须显式失败，不能静默返回空
    if total_pages >= _SCAN_MIN_PAGES and text_chars < _SCAN_CHARS_PER_PAGE * total_pages:
        raise DocumentParseError(
            f"PDF 几乎无文本层（{text_chars} 字符 / {total_pages} 页），很可能是扫描件。"
            f"本项目不做 OCR，请提供文字可选中的文本版 PDF。"
        )

    # ← 自检必须在剔除页眉页脚【之前】做 —— 页码正是接下来要被剔掉的对象
    page_audit = _audit_pages(pages)

    pages = _strip_headers_footers(pages)

    blocks = [
        Block(text=line, page=page_no, kind="paragraph")
        for page_no, lines in enumerate(pages, start=1)
        for line in lines
    ]
    return blocks, {"total_pages": total_pages, "text_chars": text_chars, **page_audit}


def parse_docx(path: Path) -> tuple[list[Block], dict]:
    try:
        document = DocxDocument(str(path))
    except Exception as exc:
        raise DocumentParseError(f"无法打开 DOCX：{path.name}（{exc}）") from exc

    blocks: list[Block] = []
    for para in document.paragraphs:
        text = _normalize(para.text)
        if not text:
            continue

        paragraph_style = para.style
        style = (paragraph_style.name or "").lower() if paragraph_style is not None else ""
        is_heading = style.startswith("heading") or style.startswith("标题")
        blocks.append(Block(text=text, page=None, kind="heading" if is_heading else "paragraph"))

    # 表格拍平成文本行。" | " 分隔保留了一点结构，但不做单元格级还原
    for table in document.tables:
        for row in table.rows:
            cells = [_normalize(c.text) for c in row.cells]
            cells = [c for c in cells if c]
            if cells:
                blocks.append(Block(text=" | ".join(cells), page=None, kind="table"))

    # 注意：DOCX 没有"页"的概念，page 只能是 None
    return blocks, {"paragraphs": len(document.paragraphs), "tables": len(document.tables)}


def parse_markdown(path: Path) -> tuple[list[Block], dict]:
    try:
        raw = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise DocumentParseError(f"Markdown 不是 UTF-8 编码：{path.name}。请转成 UTF-8 后重试。") from exc

    blocks: list[Block] = []
    for line in raw.splitlines():
        stripped = _strip_markdown(_normalize(line))
        if not stripped:
            continue
        blocks.append(Block(text=stripped, page=None, kind="paragraph"))

    return blocks, {"lines": len(blocks)}


_PARSERS: dict[str, tuple[str, Callable[[Path], tuple[list[Block], dict]]]] = {
    ".pdf": ("pdf", parse_pdf),
    ".docx": ("docx", parse_docx),
    ".md": ("md", parse_markdown),
    ".markdown": ("md", parse_markdown),
    ".txt": ("md", parse_markdown),
}

SUPPORTED_SUFFIXES = tuple(sorted(_PARSERS))


def compute_doc_id(path: Path, size: int = 16) -> str:
    """按文件内容哈希生成 doc_id。

    ← 关键：确定性。同一份合同重复上传得到相同 doc_id，
    配合 Milvus upsert 天然幂等（W3 Day 19）。用 uuid4 则每传一次多一份副本。
    """
    digest = hashlib.sha1()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()[:size]


def parse_document(path: str | Path, doc_id: str | None = None) -> ParsedDocument:
    """解析入口。

    Args:
        path: 文档路径（相对路径按当前工作目录解析——这是用户输入，不是配置）
        doc_id: 缺省时按文件内容哈希生成
    """
    path = Path(path)
    if not path.is_file():
        raise DocumentParseError(f"文件不存在：{path}")

    suffix = path.suffix.lower()
    if suffix not in _PARSERS:
        raise DocumentParseError(f"不支持的格式 {suffix!r}，当前支持：{list(SUPPORTED_SUFFIXES)}")

    source_type, parser = _PARSERS[suffix]
    blocks, meta = parser(path)

    if not blocks:
        raise DocumentParseError(f"解析结果为空：{path.name}")

    raw_text = "\n".join(b.text for b in blocks)
    meta["char_count"] = len(raw_text)

    logger.info(f"解析完成 {path.name} | 类型={source_type} 块={len(blocks)} 字符={len(raw_text)}")

    return ParsedDocument(
        doc_id=doc_id or compute_doc_id(path),
        filename=path.name,
        source_type=source_type,  # type: ignore[arg-type]
        blocks=blocks,
        raw_text=raw_text,
        meta=meta,
    )
